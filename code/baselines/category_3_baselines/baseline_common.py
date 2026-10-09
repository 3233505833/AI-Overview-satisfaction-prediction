import argparse
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TRAIN_PATH = PROJECT_ROOT / "train_merged_full_fields.jsonl"
DEFAULT_TEST_PATH = PROJECT_ROOT / "test_merged_full_fields.jsonl"
DEFAULT_USEFULNESS_PATH = PROJECT_ROOT / "query_output_train_and_test.jsonl"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "outputs"

FEATURE_NAMES = [
    "aio_sentence_usefulness_min",
    "aio_sentence_usefulness_mean",
    "aio_sentence_usefulness_max",
    "aio_examination_count",
    "aio_examination_dwell_sum",
    "aio_examination_dwell_min",
    "aio_examination_dwell_max",
    "aio_examination_dwell_mean",
    "trad_examination_count",
    "trad_click_count",
    "trad_serp_dwell_sum",
    "trad_serp_dwell_min",
    "trad_serp_dwell_max",
    "trad_serp_dwell_mean",
    "trad_content_dwell_sum",
    "trad_content_dwell_min",
    "trad_content_dwell_max",
    "trad_content_dwell_mean",
    "exposed_rank_min",
    "exposed_rank_max",
    "exposed_rank_mean",
    "clicked_rank_min",
    "clicked_rank_max",
    "clicked_rank_mean",
    "has_next_query",
    "reformulation_jaccard",
]


@dataclass
class Example:
    search_id: str
    query: str
    label: float
    current_index: int
    events: List[Dict[str, Any]]
    next_query: str
    has_next_query: bool
    search_results: List[Dict[str, Any]]


def build_arg_parser(description: str) -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument("--train", default=str(DEFAULT_TRAIN_PATH))
    parser.add_argument("--test", default=str(DEFAULT_TEST_PATH))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--output", default="")
    parser.add_argument("--usefulness-jsonl", default=str(DEFAULT_USEFULNESS_PATH))
    parser.add_argument("--seed", type=int, default=13)
    return parser


def read_jsonl(path: str | Path) -> Iterable[Dict[str, Any]]:
    with Path(path).open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{path}:{line_no} is not valid JSON: {exc}") from exc
            if isinstance(value, dict):
                yield value


def remove_ad_text(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace("广告", "")
    if isinstance(value, list):
        return [remove_ad_text(item) for item in value]
    if isinstance(value, dict):
        return {key: remove_ad_text(item) for key, item in value.items()}
    return value


def as_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "-" or value == "":
        return default
    try:
        result = float(str(value).replace("秒", "").strip())
    except Exception:
        return default
    return result if math.isfinite(result) else default


def as_optional_float(value: Any) -> Optional[float]:
    if value is None or value == "-" or value == "":
        return None
    try:
        result = float(str(value).replace("秒", "").strip())
    except Exception:
        return None
    return result if math.isfinite(result) else None


def as_int(value: Any, default: int = 0) -> int:
    if value is None or value == "-" or value == "":
        return default
    try:
        return int(float(str(value).strip()))
    except Exception:
        return default


def normalize_text(value: Any) -> str:
    return str(value or "").strip()


def normalize_search_results(search_results: Any) -> List[Dict[str, Any]]:
    while isinstance(search_results, list) and len(search_results) == 1 and isinstance(search_results[0], list):
        search_results = search_results[0]
    if not isinstance(search_results, list):
        return []
    return [item for item in search_results if isinstance(item, dict)]


def find_current_index(data: Dict[str, Any]) -> int:
    query_chain = data.get("query链路", {}) if isinstance(data.get("query链路"), dict) else {}
    query_list = list(query_chain.get("query_list", []) or [])
    search_id_list = [str(item) for item in (data.get("search_id_list", []) or [])]
    search_id = str(data.get("search_id", ""))
    if search_id and search_id in search_id_list:
        return search_id_list.index(search_id)
    current_query = data.get("用户当前查询query", "")
    matching = [idx for idx, query in enumerate(query_list) if query == current_query]
    if not matching:
        return -1
    search_pages = data.get("搜索结果页内容", []) or []
    for idx in matching:
        if idx < len(search_pages) and normalize_search_results(search_pages[idx]):
            return idx
    return matching[0]


def current_search_results(data: Dict[str, Any], current_index: int) -> List[Dict[str, Any]]:
    pages = data.get("搜索结果页内容", []) or []
    if 0 <= current_index < len(pages):
        return normalize_search_results(pages[current_index])
    return []


def select_ai_result(search_results: Sequence[Dict[str, Any]]) -> Tuple[Optional[Dict[str, Any]], str]:
    for result in search_results:
        if str(result.get("相对位置", "")) != "1":
            continue
        source = normalize_text(result.get("结果来源", ""))
        summary = normalize_text(result.get("结果摘要", ""))
        if (
            source in {"某平台AI", "某平台AI+", "内容由AI生成"}
            or result.get("是否AI摘要", "") == "是"
            or "AI" in summary
            or "智能" in summary
            or "大模型" in summary
        ):
            return result, "valid"
        return result, "invalid"
    return None, "no_position_1"


def current_pv_behavior(data: Dict[str, Any], current_index: int) -> List[Dict[str, Any]]:
    current_query = data.get("用户当前查询query", "")
    pv_list = data.get("pv行为", []) or []
    if 0 <= current_index < len(pv_list):
        item = pv_list[current_index]
        if isinstance(item, dict) and item.get("query") == current_query:
            return item.get("pv行为", []) or []
    for item in pv_list:
        if isinstance(item, dict) and item.get("query") == current_query:
            return item.get("pv行为", []) or []
    return []


def next_query_for_index(data: Dict[str, Any], current_index: int) -> str:
    query_chain = data.get("query链路", {}) if isinstance(data.get("query链路"), dict) else {}
    query_list = list(query_chain.get("query_list", []) or [])
    if 0 <= current_index + 1 < len(query_list):
        return normalize_text(query_list[current_index + 1])
    return ""


def value_at(values: Any, index: int, default: Any = "-") -> Any:
    if isinstance(values, list) and index < len(values):
        return values[index]
    return default


def flatten_pv_events(pv_behavior: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    events: List[Dict[str, Any]] = []
    keys = [
        "行为时间",
        "一级行为",
        "二级行为",
        "url",
        "页码",
        "结果类型",
        "相对位置",
        "A页曝光时长",
        "C页停留时长",
        "自动播时长",
        "点击次数",
    ]
    for group in pv_behavior:
        if not isinstance(group, dict):
            continue
        n = max((len(group.get(key, []) or []) for key in keys if isinstance(group.get(key, []), list)), default=0)
        for idx in range(n):
            action = normalize_text(value_at(group.get("一级行为"), idx))
            sub_action = normalize_text(value_at(group.get("二级行为"), idx))
            position = value_at(group.get("相对位置"), idx)
            events.append(
                {
                    "time": value_at(group.get("行为时间"), idx),
                    "action": action,
                    "sub_action": sub_action,
                    "url": value_at(group.get("url"), idx),
                    "page": as_int(value_at(group.get("页码"), idx), 0),
                    "result_type": normalize_text(value_at(group.get("结果类型"), idx)),
                    "position": as_int(position, -1),
                    "position_raw": position,
                    "a_dwell": as_optional_float(value_at(group.get("A页曝光时长"), idx)),
                    "c_dwell": as_optional_float(value_at(group.get("C页停留时长"), idx)),
                    "autoplay": as_optional_float(value_at(group.get("自动播时长"), idx)),
                    "click_count_field": as_int(value_at(group.get("点击次数"), idx), 0),
                }
            )
    return events


def is_exposure_event(event: Dict[str, Any]) -> bool:
    return "曝光" in normalize_text(event.get("action", ""))


def is_click_event(event: Dict[str, Any]) -> bool:
    text = normalize_text(event.get("action", "")) + normalize_text(event.get("sub_action", ""))
    return "点击" in text or event.get("c_dwell") is not None or as_int(event.get("click_count_field"), 0) > 0


def load_examples(path: str | Path, require_valid_aio: bool = True) -> List[Example]:
    examples: List[Example] = []
    for raw in read_jsonl(path):
        label = as_optional_float(raw.get("用户打分"))
        if label is None or label < 1 or label > 5:
            continue
        if "有广告/提示付费" in (raw.get("用户多选", []) or []):
            continue
        data = remove_ad_text(raw)
        current_index = find_current_index(data)
        if current_index < 0:
            continue
        search_results = current_search_results(data, current_index)
        top1, status = select_ai_result(search_results)
        if require_valid_aio and status != "valid":
            continue
        if top1 is None:
            continue
        query = normalize_text(data.get("用户当前查询query", ""))
        examples.append(
            Example(
                search_id=normalize_text(data.get("search_id", "")),
                query=query,
                label=float(label),
                current_index=current_index,
                events=flatten_pv_events(current_pv_behavior(data, current_index)),
                next_query=next_query_for_index(data, current_index),
                has_next_query=bool(next_query_for_index(data, current_index)),
                search_results=list(search_results),
            )
        )
    return examples


def char_jaccard(left: str, right: str) -> float:
    left_set = {ch for ch in normalize_text(left) if not ch.isspace()}
    right_set = {ch for ch in normalize_text(right) if not ch.isspace()}
    if not left_set and not right_set:
        return 1.0
    union = left_set | right_set
    return len(left_set & right_set) / len(union) if union else 0.0


def duration_stats(values: Sequence[float]) -> Dict[str, float]:
    if not values:
        return {"sum": 0.0, "min": 0.0, "max": 0.0, "mean": 0.0}
    arr = np.array(values, dtype=np.float64)
    return {
        "sum": float(arr.sum()),
        "min": float(arr.min()),
        "max": float(arr.max()),
        "mean": float(arr.mean()),
    }


def rank_stats(values: Sequence[int]) -> Dict[str, float]:
    if not values:
        return {"min": 0.0, "max": 0.0, "mean": 0.0}
    arr = np.array(values, dtype=np.float64)
    return {"min": float(arr.min()), "max": float(arr.max()), "mean": float(arr.mean())}


def load_usefulness_scores(path: str | Path) -> Dict[str, List[float]]:
    if not path:
        path = DEFAULT_USEFULNESS_PATH
    path = Path(path)
    if not path.exists():
        return {}
    result: Dict[str, List[float]] = {}
    for item in read_jsonl(path):
        search_id = normalize_text(item.get("search_id", ""))
        returned = item.get("gpt_returned_result", {})
        if not search_id or not isinstance(returned, dict):
            continue
        scores = []
        for signal in returned.get("content-level signals", []) or []:
            if not isinstance(signal, dict):
                continue
            score = as_optional_float(signal.get("usefulness_score"))
            if score is not None:
                scores.append(float(score))
        if scores:
            result[search_id] = scores
    return result


def example_usefulness_scores(example: Example, usefulness_by_id: Optional[Dict[str, List[float]]] = None) -> List[float]:
    if usefulness_by_id and example.search_id in usefulness_by_id:
        return usefulness_by_id[example.search_id]
    return [3.0]


def behavior_feature_dict(example: Example, usefulness_by_id: Optional[Dict[str, List[float]]] = None) -> Dict[str, float]:
    values = {name: 0.0 for name in FEATURE_NAMES}
    aio_dwell: List[float] = []
    trad_serp_dwell: List[float] = []
    trad_content_dwell: List[float] = []
    exposed_ranks: List[int] = []
    clicked_ranks: List[int] = []
    for event in example.events:
        position = as_int(event.get("position"), -1)
        is_aio = position == 1
        is_trad = position > 1
        exposure = is_exposure_event(event)
        click = is_click_event(event)
        a_dwell = event.get("a_dwell")
        c_dwell = event.get("c_dwell")
        if is_aio and exposure:
            values["aio_examination_count"] += 1.0
            if a_dwell is not None:
                aio_dwell.append(float(a_dwell))
        if is_trad and exposure:
            values["trad_examination_count"] += 1.0
            exposed_ranks.append(position)
            if a_dwell is not None:
                trad_serp_dwell.append(float(a_dwell))
        if is_trad and click:
            values["trad_click_count"] += 1.0
            clicked_ranks.append(position)
            if c_dwell is not None:
                trad_content_dwell.append(float(c_dwell))
    for key, stats in {
        "aio_examination_dwell": duration_stats(aio_dwell),
        "trad_serp_dwell": duration_stats(trad_serp_dwell),
        "trad_content_dwell": duration_stats(trad_content_dwell),
    }.items():
        for stat_name, value in stats.items():
            values[f"{key}_{stat_name}"] = value
    for key, stats in {"exposed_rank": rank_stats(exposed_ranks), "clicked_rank": rank_stats(clicked_ranks)}.items():
        for stat_name, value in stats.items():
            values[f"{key}_{stat_name}"] = value
    scores = np.array(example_usefulness_scores(example, usefulness_by_id), dtype=np.float64)
    values["aio_sentence_usefulness_min"] = float(scores.min()) if len(scores) else 3.0
    values["aio_sentence_usefulness_mean"] = float(scores.mean()) if len(scores) else 3.0
    values["aio_sentence_usefulness_max"] = float(scores.max()) if len(scores) else 3.0
    values["has_next_query"] = 1.0 if example.has_next_query else 0.0
    values["reformulation_jaccard"] = char_jaccard(example.query, example.next_query) if example.has_next_query else 0.0
    return values


def labels_for_examples(examples: Sequence[Example]) -> np.ndarray:
    return np.array([example.label for example in examples], dtype=np.float64)


def behavior_feature_matrix(
    examples: Sequence[Example],
    feature_names: Sequence[str] = FEATURE_NAMES,
    usefulness_by_id: Optional[Dict[str, List[float]]] = None,
) -> np.ndarray:
    rows = []
    for example in examples:
        features = behavior_feature_dict(example, usefulness_by_id)
        rows.append([features.get(name, 0.0) for name in feature_names])
    return np.array(rows, dtype=np.float64)


def log_transform_duration_columns(matrix: np.ndarray, feature_names: Sequence[str]) -> np.ndarray:
    matrix = np.array(matrix, dtype=np.float64, copy=True)
    for idx, name in enumerate(feature_names):
        if "dwell" in name:
            matrix[:, idx] = np.log1p(np.maximum(matrix[:, idx], 0.0))
    return matrix


def fit_standardizer(matrix: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    mean = matrix.mean(axis=0)
    std = matrix.std(axis=0)
    std[std < 1e-8] = 1.0
    return mean, std


def apply_standardizer(matrix: np.ndarray, mean: np.ndarray, std: np.ndarray) -> np.ndarray:
    return (matrix - mean) / std


def prepare_tabular_features(
    train_examples: Sequence[Example],
    test_examples: Sequence[Example],
    feature_names: Sequence[str] = FEATURE_NAMES,
    usefulness_by_id: Optional[Dict[str, List[float]]] = None,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    x_train = log_transform_duration_columns(behavior_feature_matrix(train_examples, feature_names, usefulness_by_id), feature_names)
    x_test = log_transform_duration_columns(behavior_feature_matrix(test_examples, feature_names, usefulness_by_id), feature_names)
    mean, std = fit_standardizer(x_train)
    return apply_standardizer(x_train, mean, std), apply_standardizer(x_test, mean, std), mean, std


def sequence_event_vector(event: Dict[str, Any], order: int, max_len: int, usefulness_mean: float) -> List[float]:
    position = as_int(event.get("position"), -1)
    exposure = 1.0 if is_exposure_event(event) else 0.0
    click = 1.0 if is_click_event(event) else 0.0
    a_dwell = math.log1p(max(as_float(event.get("a_dwell")), 0.0))
    c_dwell = math.log1p(max(as_float(event.get("c_dwell")), 0.0))
    autoplay = math.log1p(max(as_float(event.get("autoplay")), 0.0))
    return [
        exposure,
        click,
        1.0 if position == 1 else 0.0,
        1.0 if position > 1 else 0.0,
        max(position, 0) / 10.0,
        a_dwell,
        c_dwell,
        autoplay,
        order / max(max_len - 1, 1),
        usefulness_mean / 5.0,
    ]


def sequence_tensor(
    examples: Sequence[Example],
    max_len: int = 64,
    usefulness_by_id: Optional[Dict[str, List[float]]] = None,
) -> Tuple[np.ndarray, np.ndarray]:
    feature_dim = 10
    x = np.zeros((len(examples), max_len, feature_dim), dtype=np.float64)
    mask = np.zeros((len(examples), max_len), dtype=np.float64)
    for row_idx, example in enumerate(examples):
        scores = example_usefulness_scores(example, usefulness_by_id)
        usefulness_mean = float(np.mean(np.array(scores, dtype=np.float64))) if scores else 3.0
        for event_idx, event in enumerate(example.events[:max_len]):
            x[row_idx, event_idx, :] = sequence_event_vector(event, event_idx, max_len, usefulness_mean)
            mask[row_idx, event_idx] = 1.0
    return x, mask


def clip_predictions(predictions: Sequence[float]) -> np.ndarray:
    return np.clip(np.array(predictions, dtype=np.float64), 1.0, 5.0)


def xlsx_safe_sheet_name(name: str) -> str:
    clean = re.sub(r"[\[\]\*:/\\?]", "_", name)
    return clean[:31] or "baseline"


def default_output_path(output_dir: str | Path, baseline_name: str) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    file_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", baseline_name).strip("_")
    return output_dir / f"{file_name}.xlsx"


def save_predictions_excel(
    predictions: Sequence[float],
    truths: Sequence[float],
    output_path: str | Path,
    sheet_name: str,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = xlsx_safe_sheet_name(sheet_name)
    worksheet.append(["满意度预测", "用户打分"])
    for prediction, truth in zip(predictions, truths):
        worksheet.append([float(prediction), float(truth)])
    workbook.save(output_path)
    return output_path


def save_combined_workbook(results: Sequence[Dict[str, Any]], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for result in results:
        worksheet = workbook.create_sheet(xlsx_safe_sheet_name(result["name"]))
        worksheet.append(["满意度预测", "用户打分"])
        for prediction, truth in zip(result["predictions"], result["truths"]):
            worksheet.append([float(prediction), float(truth)])
    workbook.save(output_path)
    return output_path


def print_run_summary(name: str, train_examples: Sequence[Example], test_examples: Sequence[Example], output_path: Path) -> None:
    print(f"{name}: train={len(train_examples)}, test={len(test_examples)}, output={output_path}")
